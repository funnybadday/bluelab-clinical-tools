import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT_DIR))

from scripts.phase2.fill_table import fill_table  # noqa: E402


class TwoGroupTteTemplateTests(unittest.TestCase):
    def test_two_endpoints_generate_one_complete_tte_table(self):
        """Prevent F_G2_TTE from falling back to an unfilled static shell."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            data_path = temp_path / "tte-data.json"
            output_path = temp_path / "tte.docx"
            data_path.write_text(
                json.dumps(
                    {
                        "projects": [
                            {"name": "无进展生存期", "unit": "月"},
                            {"name": "总生存期"},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = fill_table("F_G2_TTE", str(data_path), str(output_path))

            self.assertEqual(result, str(output_path))
            document = Document(output_path)
            self.assertEqual(len(document.tables), 1)

            table = document.tables[0]
            self.assertEqual(len(table.rows), 45)
            self.assertEqual(
                [cell.text for cell in table.rows[0].cells],
                ["项目", "指标", "试验组\n(N=n1)", "对照组\n(N=n1)"],
            )
            self.assertEqual(table.cell(1, 0).text, "无进展生存期")
            self.assertEqual(table.cell(5, 1).text, "无进展生存期（月）")
            self.assertEqual(table.cell(14, 1).text, "Logrank检验")
            self.assertEqual(table.cell(15, 2).text, "x.xxx")
            self.assertEqual(table.cell(15, 3).text, "")
            self.assertEqual(table.cell(23, 0).text, "总生存期")
            self.assertEqual(table.cell(27, 1).text, "总生存期（天）")


class SingleGroupTteTemplateTests(unittest.TestCase):
    def test_two_endpoints_generate_one_complete_single_group_tte_table(self):
        """Prevent F_G1_TTE from falling back to its unfilled three-column shell."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            data_path = temp_path / "tte-data.json"
            output_path = temp_path / "tte.docx"
            data_path.write_text(
                json.dumps(
                    {
                        "projects": [
                            {"name": "无进展生存期", "unit": "月"},
                            {"name": "总生存期"},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = fill_table("F_G1_TTE", str(data_path), str(output_path))

            self.assertEqual(result, str(output_path))
            document = Document(output_path)
            self.assertEqual(len(document.tables), 1)

            table = document.tables[0]
            self.assertEqual(len(table.rows), 25)
            self.assertEqual(
                [cell.text for cell in table.rows[0].cells],
                ["项目", "指标", "结果\n(N=n1)"],
            )
            self.assertEqual(table.cell(1, 0).text, "无进展生存期")
            self.assertEqual(table.cell(1, 2).text, "n3 (n4)")
            self.assertEqual(table.cell(5, 1).text, "无进展生存期（月）")
            self.assertEqual(table.cell(6, 2).text, "x.xx")
            self.assertEqual(table.cell(13, 0).text, "总生存期")
            self.assertEqual(table.cell(17, 1).text, "总生存期（天）")
            self.assertNotIn("Logrank检验", "\n".join(cell.text for row in table.rows for cell in row.cells))
            self.assertNotIn("Cox回归", "\n".join(cell.text for row in table.rows for cell in row.cells))

            column_widths = [column.width.twips for column in table.columns]
            self.assertLess(column_widths[0], column_widths[2])
            self.assertGreater(column_widths[1], column_widths[0])
            self.assertGreater(column_widths[1], column_widths[2])
            self.assertTrue(table.autofit)


if __name__ == "__main__":
    unittest.main()
