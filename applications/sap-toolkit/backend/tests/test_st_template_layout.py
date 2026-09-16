import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT


ROOT_DIR = Path(__file__).resolve().parents[2]
PHASE2_DIR = ROOT_DIR / "scripts" / "phase2"
sys.path.insert(0, str(PHASE2_DIR))

from extract_template import (  # noqa: E402
    extract_multiple_tables,
    extract_table_by_code,
    load_template_library,
)
from format_word_tables import process_document  # noqa: E402


ST_TEMPLATE_CODES = ("D_G1_ST", "D_G2_ST", "D_G3_ST")


def table_layout(table):
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    layout = table_properties.first_child_found_in("w:tblLayout")
    grid_width = sum(int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst)
    return (
        table_width.get(qn("w:w")),
        table_width.get(qn("w:type")),
        layout.get(qn("w:type")),
        grid_width,
    )


class StTemplateLayoutTests(unittest.TestCase):
    def test_st_templates_extract_on_landscape_pages_with_their_master_layout(self):
        master, code_to_table_index = load_template_library()
        with tempfile.TemporaryDirectory() as temp_dir:
            for code in ST_TEMPLATE_CODES:
                with self.subTest(code=code):
                    output_path = Path(temp_dir) / f"{code}.docx"
                    extract_table_by_code(code, output_path=output_path)

                    document = Document(output_path)
                    source_layout = table_layout(
                        master.tables[code_to_table_index[code]]
                    )

                    self.assertEqual(document.sections[0].orientation, WD_ORIENT.LANDSCAPE)
                    self.assertGreater(
                        document.sections[0].page_width,
                        document.sections[0].page_height,
                    )
                    self.assertEqual(table_layout(document.tables[0]), source_layout)

    def test_two_group_st_master_uses_trial_and_control_headers(self):
        master, code_to_table_index = load_template_library()
        table = master.tables[code_to_table_index["D_G2_ST"]]

        self.assertEqual(
            [cell.text for cell in table.rows[1].cells],
            [
                "中心编号",
                "试验组",
                "对照组",
                "合计",
                "试验组",
                "对照组",
                "合计",
                "试验组",
                "对照组",
                "合计",
            ],
        )

    def test_multiple_extraction_keeps_landscape_when_it_includes_st_templates(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "st-tables.docx"
            extract_multiple_tables(["D_G1_ST", "D_G2_ST"], output_path=output_path)

            document = Document(output_path)
            self.assertEqual(document.sections[0].orientation, WD_ORIENT.LANDSCAPE)
            self.assertGreater(
                document.sections[0].page_width,
                document.sections[0].page_height,
            )

    def test_formatting_population_by_center_tables_uses_full_page_width(self):
        """Prevent D_G*_ST tables from retaining a narrower master width."""
        with tempfile.TemporaryDirectory() as temp_dir:
            for code in ST_TEMPLATE_CODES:
                with self.subTest(code=code):
                    source_path = Path(temp_dir) / f"{code}.docx"
                    extract_table_by_code(code, output_path=source_path)

                    formatted_path = process_document(source_path)
                    table_properties = Document(formatted_path).tables[0]._tbl.tblPr
                    table_width = table_properties.first_child_found_in("w:tblW")

                    self.assertEqual(table_width.get(qn("w:type")), "pct")
                    self.assertEqual(table_width.get(qn("w:w")), "5000")


if __name__ == "__main__":
    unittest.main()
