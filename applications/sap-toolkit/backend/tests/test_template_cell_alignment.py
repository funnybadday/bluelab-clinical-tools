import importlib.util
import inspect
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT_DIR = Path(__file__).resolve().parents[2]
TEMPLATE_ROOT = ROOT_DIR / "scripts" / "phase2" / "模板代码"
CENTERED_TEMPLATE_CODES = (
    "D_G1_EN",
    "D_G2_EN",
    "D_G3_EN",
    "E_G1_FREQ1",
    "E_G2_FREQ1",
    "E_G3_FREQ1",
    "E_G3_FREQ3",
    "F_G1_QUALCOM_N_N",
    "F_G2_CMH",
    "F_G2_LOGISTIC",
    "F_G2_OR",
    "F_G2_RR",
    "F_G2_RR_EQU",
    "F_G2_RR_NIN",
    "F_G2_RR_SUP",
)


def load_module(module_path: Path):
    spec = importlib.util.spec_from_file_location(module_path.parent.name, module_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def render_template(template_code: str) -> Document:
    template_dir = TEMPLATE_ROOT / template_code
    with (template_dir / f"{template_code}.json").open(encoding="utf-8") as f:
        semantic = json.load(f)
    with (template_dir / "填充数据.json").open(encoding="utf-8") as f:
        data = json.load(f)

    fill_module = load_module(template_dir / "fill_template.py")
    fill_args = [semantic, data["projects"]]
    if "visits" in inspect.signature(fill_module.fill_template).parameters:
        fill_args.append(data.get("visits", []))
    filled = fill_module.fill_template(*fill_args)

    with tempfile.TemporaryDirectory() as temp_dir:
        output_path = Path(temp_dir) / "table.docx"
        gen_docx_module = load_module(template_dir / "gen_docx.py")
        gen_docx_module.generate_docx(filled, output_path)
        return Document(output_path)


class TemplateCellAlignmentTests(unittest.TestCase):
    def test_centered_templates_render_all_nonempty_cell_text_left_aligned(self):
        for template_code in CENTERED_TEMPLATE_CODES:
            with self.subTest(template_code=template_code):
                document = render_template(template_code)
                alignments = [
                    paragraph.alignment
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                    for paragraph in cell.paragraphs
                    if paragraph.text
                ]

                self.assertTrue(alignments)
                self.assertTrue(
                    all(alignment == WD_ALIGN_PARAGRAPH.LEFT for alignment in alignments),
                    f"{template_code} still contains non-left-aligned cell content",
                )


if __name__ == "__main__":
    unittest.main()
