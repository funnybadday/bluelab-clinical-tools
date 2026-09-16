import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT_DIR))

from scripts.phase2.fill_table import fill_table  # noqa: E402


class TwoGroupMmrmTemplateTests(unittest.TestCase):
    def _fill(self, template_code, payload):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        temp_path = Path(temp_dir.name)
        data_path = temp_path / f"{template_code}.json"
        output_path = temp_path / f"{template_code}.docx"
        data_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

        result = fill_table(template_code, str(data_path), str(output_path))
        self.assertEqual(result, str(output_path))
        document = Document(output_path)
        self.assertEqual(len(document.tables), 1)
        return document.tables[0]

    def test_indicator_mmrm_templates_generate_one_two_group_table_per_input_indicator(self):
        """Prevent the three non-visit G2 MMRM shells from remaining unfillable."""
        payload = {
            "projects": [
                {"name": "症状评分", "unit": "分"},
                {"name": "生活质量评分", "unit": "分"},
            ]
        }
        expected_layouts = {
            "F_G2_MMRM_P1": (11, 6),
            "F_G2_MMRM_P2": (15, 8),
            "F_G2_MMRM_02_P1": (15, 8),
        }

        for template_code, (row_count, second_start) in expected_layouts.items():
            with self.subTest(template_code=template_code):
                table = self._fill(template_code, payload)
                self.assertEqual(len(table.rows), row_count)
                self.assertEqual(len(table.columns), 4)
                self.assertEqual(table.cell(1, 0).text, "症状评分 (分)")
                self.assertEqual(table.cell(second_start, 0).text, "生活质量评分 (分)")

        comparison_table = self._fill("F_G2_MMRM_P2", payload)
        self.assertEqual(
            [cell.text for cell in comparison_table.rows[0].cells],
            ["项目", "指标", "试验组\n(N=n1)", "对照组\n(N=n1)"],
        )
        self.assertEqual(comparison_table.cell(4, 1).text, "两组最小二乘均数差值")
        self.assertEqual(comparison_table.cell(4, 2).text, "x.xx")
        self.assertEqual(comparison_table.cell(4, 3).text, "")

    def test_visit_mmrm_template_uses_the_supplied_real_visit_names(self):
        """Prevent MMRM visit output from substituting the generic fallback schedule."""
        table = self._fill(
            "F_G2_MMRM_02_P2",
            {
                "projects": [{"name": "症状评分", "unit": "分"}],
                "visits": ["基线", "治疗后4周", "治疗后8周"],
            },
        )

        self.assertEqual(len(table.rows), 25)
        self.assertEqual(table.cell(1, 0).text, "基线")
        self.assertEqual(table.cell(9, 0).text, "治疗后4周")
        self.assertEqual(table.cell(17, 0).text, "治疗后8周")
        self.assertEqual(table.cell(4, 1).text, "两组最小二乘均数差值")
        self.assertEqual(table.cell(4, 2).text, "x.xx")
        self.assertEqual(table.cell(4, 3).text, "")

    def test_visit_mmrm_template_keeps_the_original_visit_placeholder_when_no_visits_are_available(self):
        """Prevent missing visit data from becoming a fictional default schedule."""
        table = self._fill(
            "F_G2_MMRM_02_P2",
            {"projects": [{"name": "症状评分", "unit": "分"}]},
        )

        self.assertEqual(len(table.rows), 9)
        self.assertEqual(table.cell(1, 0).text, "XX访视")


class MmrmVisitEnrichmentTests(unittest.TestCase):
    def _create_project(self, existing_visits=None):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        output_dir = Path(temp_dir.name)
        info_dir = output_dir / "05_表格信息"
        content_dir = output_dir / "02_内容提取"
        info_dir.mkdir(parents=True)
        content_dir.mkdir(parents=True)

        table_name = "症状评分-重复测量的混合效应模型估计情况（FAS）"
        table_info = {
            "table_name": table_name,
            "endpoint": {"name": "症状评分", "unit": "分"},
            "projects": [{"name": "症状评分", "unit": "分"}],
        }
        if existing_visits is not None:
            table_info["visits"] = existing_visits
        info_path = info_dir / f"{table_name}.json"
        info_path.write_text(json.dumps(table_info, ensure_ascii=False), encoding="utf-8")
        (content_dir / "试验流程.md").write_text(
            "| 时间点 | 基线 | 治疗后4周 | 治疗后8周 |\n| 症状评分 | X | X | X |",
            encoding="utf-8",
        )
        (output_dir / "模板代码结果.json").write_text(
            json.dumps(
                {
                    "tables": [
                        {"name": table_name, "template_code": "F_G2_MMRM_02_P2"}
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return output_dir, info_path

    def test_mmrm_visit_enrichment_writes_extracted_endpoint_visits_to_only_the_target_table(self):
        """Prevent F_G2_MMRM_02_P2 from reaching its renderer without real visits."""
        from scripts.phase2.enrich_mmrm_visits import enrich_mmrm_visits

        output_dir, info_path = self._create_project()

        result = enrich_mmrm_visits(
            output_dir,
            visit_extractor=lambda flow, endpoint: ["基线", "治疗后4周", "治疗后8周"],
        )

        self.assertEqual(result, {"updated": 1, "preserved": 0, "unavailable": 0})
        saved = json.loads(info_path.read_text(encoding="utf-8"))
        self.assertEqual(saved["visits"], ["基线", "治疗后4周", "治疗后8周"])

    def test_mmrm_visit_enrichment_preserves_manually_supplied_visits(self):
        """Prevent enrichment from overwriting a user-provided endpoint schedule."""
        from scripts.phase2.enrich_mmrm_visits import enrich_mmrm_visits

        output_dir, info_path = self._create_project(["基线", "治疗后12周"])

        result = enrich_mmrm_visits(
            output_dir,
            visit_extractor=lambda flow, endpoint: self.fail("manual visits must not call extraction"),
        )

        self.assertEqual(result, {"updated": 0, "preserved": 1, "unavailable": 0})
        saved = json.loads(info_path.read_text(encoding="utf-8"))
        self.assertEqual(saved["visits"], ["基线", "治疗后12周"])


if __name__ == "__main__":
    unittest.main()
