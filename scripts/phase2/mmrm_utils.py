"""Shared row expansion and Word rendering helpers for G2 MMRM shells."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def project_display(project: dict) -> str:
    name = str(project.get("name") or "指标XXX")
    unit = str(project.get("unit") or "")
    return f"{name} ({unit})" if unit else name


def expand_indicator_rows(semantic: dict, projects: list[dict]) -> list[dict]:
    """Repeat every semantic MMRM row for each endpoint in one output table."""
    template_rows = [
        row for section in semantic.get("sections", []) for row in section.get("rows", [])
    ]
    rows = []

    for project in projects:
        display = project_display(project)
        for row_index, template_row in enumerate(template_rows):
            labels = template_row.get("label_values", [])
            metric = labels[1] if len(labels) > 1 else ""
            rows.append(
                {
                    "indicator": display if row_index == 0 else "",
                    "metric": metric,
                    "values": list(template_row.get("data_values", [])),
                    "applies_to": list(template_row.get("applies_to", [])),
                }
            )
    return rows


def expand_visit_rows(semantic: dict, visits: list[str] | None) -> list[dict]:
    """Repeat the MMRM visit block while retaining XX访视 when the source is absent."""
    sections = semantic.get("sections", [])
    template_rows = sections[0].get("rows", []) if sections else []
    rows = []

    for visit in visits or ["XX访视"]:
        for row_index, template_row in enumerate(template_rows):
            labels = template_row.get("label_values", [])
            metric = labels[1] if len(labels) > 1 else ""
            rows.append(
                {
                    "indicator": str(visit) if row_index == 0 else "",
                    "metric": metric,
                    "values": list(template_row.get("data_values", [])),
                    "applies_to": list(template_row.get("applies_to", [])),
                }
            )
    return rows


def _set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)

    run = paragraph.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:cs"), "Times New Roman")


def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None) -> None:
    properties = cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:tcBorders"))
    if existing is not None:
        properties.remove(existing)

    borders = OxmlElement("w:tcBorders")
    for side, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        border = OxmlElement(f"w:{side}")
        if spec is None or spec == "nil":
            border.set(qn("w:val"), "nil")
        else:
            border.set(qn("w:val"), "single")
            border.set(qn("w:color"), spec.get("color", "auto"))
            border.set(qn("w:sz"), str(spec.get("sz", 4)))
            border.set(qn("w:space"), "0")
        borders.append(border)
    properties.append(borders)


def _apply_three_line_borders(table) -> None:
    last_row = len(table.rows) - 1
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0:
                _set_cell_borders(
                    cell,
                    top={"sz": 4, "color": "auto"},
                    bottom={"sz": 4, "color": "000000"},
                    left="nil",
                    right="nil",
                )
            elif row_index == last_row:
                _set_cell_borders(
                    cell,
                    top="nil",
                    bottom={"sz": 4, "color": "auto"},
                    left="nil",
                    right="nil",
                )
            else:
                _set_cell_borders(cell, top="nil", bottom="nil", left="nil", right="nil")


def _set_table_width(table) -> None:
    table.autofit = True
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.insert(0, width)
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")


def _header_label(name: str) -> str:
    return name.replace(" (N=n1)", "\n(N=n1)")


def generate_mmrm_docx(filled_data: dict, output_path) -> None:
    """Render a complete MMRM shell into the single table consumed by merge_tables."""
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(1.91)
    section.right_margin = Cm(1.91)

    columns = filled_data.get("columns", [])
    rows = filled_data.get("rows", [])
    table = document.add_table(rows=1 + len(rows), cols=len(columns))
    _set_table_width(table)

    for index, column in enumerate(columns):
        _set_cell_text(table.cell(0, index), _header_label(column["name"]))

    for row_index, row in enumerate(rows, start=1):
        _set_cell_text(table.cell(row_index, 0), row.get("indicator", ""))
        _set_cell_text(table.cell(row_index, 1), row.get("metric", ""))
        values_by_column = {
            column_name: values_index < len(row.get("values", []))
            and row["values"][values_index]
            or ""
            for values_index, column_name in enumerate(row.get("applies_to", []))
        }
        for column_index, column in enumerate(columns[2:], start=2):
            _set_cell_text(
                table.cell(row_index, column_index),
                values_by_column.get(column["name"], ""),
            )

    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    _apply_three_line_borders(table)
    document.save(output_path)
